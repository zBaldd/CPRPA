import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options

from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

TERMO_BUSCA = "Montilla Rum"

def criar_driver():
    """Configura e inicializa o navegador."""
    options = Options()
    options.add_argument("--start-maximized") # Abre o navegador em tela cheia

    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option('useAutomationExtension', False)
    
    
    options.add_argument("--disable-blink-features=AutomationControlled")
    return webdriver.Chrome(options=options)

def buscar_produto(driver, termo):
    """Acessa o site do Mercado Livre e realiza a busca pelo termo desejado."""
    driver.get("https://www.mercadolivre.com.br")
    time.sleep(2)
    
    #Aguarda até que a barra de pesquisa esteja presente na tela
    campo = WebDriverWait(driver, 15).until(
        EC.presence_of_element_located((By.CSS_SELECTOR,
            "input.nav-search-input, input[placeholder*='Buscar'], input[type='text']"))
    )
    campo.send_keys(termo, Keys.ENTER)
    time.sleep(3)

def rolar_pagina(driver):
    """Rola a página para baixo gradualmente para carregar todos os elementos (Lazy Load)."""
    for _ in range(5):
        driver.execute_script("window.scrollBy(0, 800)")
        time.sleep(0.4)

def extrair_produtos_pagina(driver):
    """Extrai título, preço e link de cada produto listado na página atual."""
    rolar_pagina(driver)
    
    #Busca os blocos que contêm os produtos (considerando diferentes layouts do site)
    cards = (driver.find_elements(By.CSS_SELECTOR, "li.ui-search-layout__item")
          or driver.find_elements(By.CSS_SELECTOR, "div.poly-card"))
    produtos = []

    for card in cards:
        titulo = ""
        #Tenta diferentes seletores para encontrar o título do produto
        for sel in [".poly-component__title", "h2.ui-search-item__title", "h2", "a[title]"]:
            try:
                el = card.find_element(By.CSS_SELECTOR, sel)
                titulo = el.get_attribute("title") or el.text
                if titulo: break
            except: continue

        #Extrai o preço do produto tratando os centavos, se existirem
        try:
            inteiro = card.find_element(By.CSS_SELECTOR, ".andes-money-amount__fraction").text
            try:    preco = f"R$ {inteiro},{card.find_element(By.CSS_SELECTOR, '.andes-money-amount__cents').text}"
            except: preco = f"R$ {inteiro}"
        except: preco = "N/D"

        link = "N/D"
        #Tenta diferentes seletores para encontrar o link do anúncio
        for sel in ["a.poly-component__title", "a.ui-search-link", "a"]:
            try:
                link = card.find_element(By.CSS_SELECTOR, sel).get_attribute("href") or "N/D"
                if link != "N/D": break
            except: continue

        if titulo:
            produtos.append({"titulo": titulo, "preco": preco, "link": link})

    return produtos

def proxima_pagina(driver):
    """Clica no botão 'Próxima página', parando corretamente na última."""
    for sel in ["li.andes-pagination__button--next", "[class*='pagination__button--next']"]:
        try:
           
            botao = driver.find_element(By.CSS_SELECTOR, sel)
            
          
            if "disabled" in botao.get_attribute("class"):
                return False 
            
            
            link = botao.find_element(By.CSS_SELECTOR, "a")
            driver.execute_script("arguments[0].click();", link)
            time.sleep(3)
            return True
        except: 
            continue
            
    return False

def coletar_todos(driver):
    """Navega por todas as páginas disponíveis e consolida a lista de produtos."""
    todos, pagina = [], 1
    while True:
        todos.extend(extrair_produtos_pagina(driver))
        print(f"Pagina {pagina}: {len(todos)} coletados")
        
        # Interrompe o loop quando não houver mais botão de próxima página
        if not proxima_pagina(driver):
            break
        pagina += 1
    return todos

def filtro_rum(produtos):
    """Filtra a lista deixando apenas itens que contenham 'rum' no título."""
    return [p for p in produtos if "rum" in p["titulo"].lower()]

def filtro_carta_ouro(produtos):
    """Filtra a lista deixando apenas itens que contenham 'carta' e 'ouro' no título."""
    return [p for p in produtos if "carta" in p["titulo"].lower() and "ouro" in p["titulo"].lower()]

def exibir(produtos, rotulo):
    """Imprime os resultados no terminal de forma formatada."""
    print(f"\n{rotulo} [{len(produtos)} resultado(s)]")
    for i, p in enumerate(produtos, 1):
        print(f"  [{i:02d}] {p['titulo']} | {p['preco']}")
        print(f"        {p['link']}")

def salvar_xlsx(produtos, nome):
    """Gera e salva uma planilha Excel formatada com os dados extraídos."""
    wb = Workbook()
    ws = wb.active
    
    # Cria o cabeçalho
    ws.append(["preco", "nome", "link"])
    
    # Aplica estilos (cor, fonte e alinhamento) no cabeçalho
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="2D6A4F")
        cell.alignment = Alignment(horizontal="center")
        
    # Insere os dados dos produtos linha a linha
    for p in produtos:
        ws.append([p["preco"], p["titulo"], p["link"]])
        
    # Ajusta a largura das colunas
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 60
    
    # Salva o arquivo com o nome definido
    wb.save(f"{nome}.xlsx")

def main():
    """Função principal que orquestra toda a execução da automação."""
    driver = criar_driver()
    buscar_produto(driver, TERMO_BUSCA)
    todos = coletar_todos(driver)
    driver.quit() # Encerra o navegador após a coleta

    # Aplica os filtros aos dados brutos
    rum        = filtro_rum(todos)
    carta_ouro = filtro_carta_ouro(rum)

    # Exporta os resultados para planilhas separadas
    salvar_xlsx(todos,       "todos")
    salvar_xlsx(rum,         "filtro_rum")
    salvar_xlsx(carta_ouro,  "filtro_carta_ouro")

    # Exibe os resultados no console
    exibir(todos,       "TODOS")
    exibir(rum,         "FILTRO 1 - RUM")
    exibir(carta_ouro,  "FILTRO 2 - CARTA OURO")

if __name__ == "__main__":
    main()